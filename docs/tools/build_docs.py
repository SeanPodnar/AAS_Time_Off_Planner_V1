#!/usr/bin/env python3
"""Rebuild docs/files/*.docx and *.pdf from docs/PRD.md and docs/TRD.md.

Needs pandoc, LibreOffice (soffice) and python-docx. Run from anywhere:
    python3 docs/tools/build_docs.py
"""
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOCS = Path(__file__).resolve().parents[1]
OUT = DOCS / "files"
NAVY, BLUE = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x2B, 0x6C, 0xB0)


def column_widths(md: str) -> str:
    """Pandoc sizes pipe-table columns by the dashes in the separator row; size them by content."""
    lines, out, i = md.split("\n"), [], 0
    while i < len(lines):
        if lines[i].startswith("|") and i + 1 < len(lines) and re.match(r"^\|(\s*:?-+:?\s*\|)+\s*$", lines[i + 1]):
            j, rows = i, []
            while j < len(lines) and lines[j].startswith("|"):
                rows.append(lines[j])
                j += 1
            cells = [[c.strip() for c in r.strip().strip("|").split("|")] for k, r in enumerate(rows) if k != 1]
            n = len(cells[0])
            longest = [max(len(r[k].replace("**", "")) if k < len(r) else 0 for r in cells) for k in range(n)]
            widths = [max(12, len(cells[0][k]) + 1, min(longest[k], 48)) for k in range(n)]
            out += [rows[0], "|" + "|".join("-" * w for w in widths) + "|", *rows[2:]]
            i = j
        else:
            out.append(lines[i])
            i += 1
    return "\n".join(out)


def reference_docx(path: Path):
    subprocess.run(["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"], check=True)
    d = Document(str(path))
    for name in ["Normal", "Body Text", "First Paragraph", "Compact"]:
        if name in d.styles:
            d.styles[name].font.name, d.styles[name].font.size = "Calibri", Pt(10.5)
    for name, size, colour in [("Title", 24, NAVY), ("Subtitle", 13, BLUE), ("Heading 1", 16, NAVY), ("Heading 2", 13, BLUE), ("Heading 3", 11.5, NAVY)]:
        if name in d.styles:
            st = d.styles[name]
            st.font.name, st.font.size, st.font.color.rgb, st.font.bold = "Calibri", Pt(size), colour, name != "Subtitle"
    d.save(str(path))


def style_tables(path: Path):
    d = Document(str(path))
    for t in d.tables:
        borders = OxmlElement("w:tblBorders")
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "B8C7D9")
            borders.append(el)
        t._tbl.tblPr.append(borders)
        for cell in t.rows[0].cells:
            shade = OxmlElement("w:shd")
            shade.set(qn("w:val"), "clear"); shade.set(qn("w:fill"), "E8EEF6")
            cell._tc.get_or_add_tcPr().append(shade)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
    d.save(str(path))


def main():
    for tool in ["pandoc", "soffice"]:
        if not shutil.which(tool):
            sys.exit(f"{tool} is required")
    OUT.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        ref = tmp / "ref.docx"
        reference_docx(ref)
        for name in ["PRD", "TRD"]:
            src = tmp / f"{name}.md"
            src.write_text(column_widths((DOCS / f"{name}.md").read_text()))
            docx = OUT / f"AAS_Time_Off_Planner_{name}.docx"
            subprocess.run(["pandoc", str(src), "-o", str(docx), f"--reference-doc={ref}", "--columns=72"], check=True)
            style_tables(docx)
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(docx)], check=True, capture_output=True)
            print(f"built {docx.name} and PDF")


if __name__ == "__main__":
    main()
